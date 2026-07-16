from __future__ import annotations

import re
import unicodedata
from datetime import datetime
from io import BytesIO
from typing import Any
from urllib.parse import urlsplit
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

from .db import FIELD_LABELS


KST = ZoneInfo("Asia/Seoul")
NAVY = "17324D"
TEAL = "0F8B8D"
PALE_BLUE = "EAF2F8"
PALE_GREEN = "E8F5E9"
PALE_ORANGE = "FFF3E0"
PALE_GRAY = "F3F4F6"
WHITE = "FFFFFF"
TEXT = "263238"
LIGHT_BORDER = Side(style="thin", color="D9E1E8")
MAX_EXCEL_TEXT_LENGTH = 32_767
ILLEGAL_EXCEL_CHARACTERS = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")


REFERENCE_COLUMNS = [
    ("sequence", "순번"),
    ("name", "기업명"),
    ("industry", "업종"),
    ("representative", "대표자"),
    ("founded_year", "설립연도"),
    ("employee_count", "직원수"),
    ("phone", "전화번호"),
    ("address", "주소"),
    ("email", "이메일"),
    ("homepage", "홈페이지"),
    ("business_area", "사업영역"),
]

CHANGED_COMPANY_COLUMNS = [
    ("last_sync_status", "상태"),
    ("name", "기업명"),
    ("industry", "업종"),
    ("item", "세부품목"),
    ("representative", "대표자"),
    ("founded_year", "설립연도"),
    ("employee_count", "직원 수"),
    ("phone", "전화번호"),
    ("address", "주소"),
    ("email", "이메일"),
    ("homepage", "홈페이지"),
    ("business_area", "사업영역"),
    ("interest_category", "관심분야"),
    ("first_seen_at", "최초 저장"),
    ("last_seen_at", "최근 확인"),
    ("last_changed_at", "최근 변경"),
    ("source_url", "원본 페이지"),
]

STATUS_LABELS = {"new": "신규", "changed": "변경", "unchanged": "기존"}


def _clean_excel_text(value: Any) -> str:
    """Return text that is valid in an XLSX XML cell."""
    text = unicodedata.normalize("NFC", str(value))
    # The source uses vertical tabs as paragraph separators. Excel rejects them.
    text = text.replace("\x0b", "\n").replace("\x0c", "\n")
    text = ILLEGAL_EXCEL_CHARACTERS.sub("", text)
    return text[:MAX_EXCEL_TEXT_LENGTH]


def _excel_value(value: Any) -> Any:
    if value is None or not isinstance(value, str):
        return value
    return _clean_excel_text(value)


def _append_excel_row(sheet, values: list[Any]) -> None:
    sheet.append([_excel_value(value) for value in values])


def _hyperlink_target(value: Any) -> str | None:
    if not value:
        return None
    candidate = _clean_excel_text(value).strip()
    if candidate.lower().startswith("www."):
        candidate = f"https://{candidate}"
    parsed = urlsplit(candidate)
    if parsed.scheme.lower() in {"http", "https"} and parsed.netloc:
        return candidate
    return None


def _set_web_hyperlink(cell) -> None:
    target = _hyperlink_target(cell.value)
    if target:
        cell.hyperlink = target
        cell.style = "Hyperlink"


def _display_width(value: Any) -> int:
    text = "" if value is None else str(value)
    return sum(2 if unicodedata.east_asian_width(char) in {"W", "F", "A"} else 1 for char in text)


def _estimated_lines(value: Any, width: int) -> int:
    text = "" if value is None else str(value)
    logical_lines = text.splitlines() or [""]
    return sum(max(1, (_display_width(line) + width - 1) // width) for line in logical_lines)


def _style_sheet(sheet, widths: dict[int, int] | None = None) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.row_dimensions[1].height = 28
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color=WHITE, bold=True, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=Side(style="medium", color=TEAL))

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(color=TEXT, size=10)
            cell.alignment = Alignment(vertical="top")
            cell.border = Border(bottom=LIGHT_BORDER)

    for column_cells in sheet.iter_cols():
        index = column_cells[0].column
        if widths and index in widths:
            width = widths[index]
        else:
            sample = list(column_cells[:250])
            width = min(max((_display_width(cell.value) for cell in sample), default=8) + 2, 42)
        sheet.column_dimensions[column_cells[0].column_letter].width = max(8, width)


def _style_reference_sheet(sheet) -> None:
    sheet.freeze_panes = "A4"
    sheet.row_dimensions[1].height = 28
    sheet.row_dimensions[2].height = 8
    sheet.row_dimensions[3].height = 24
    sheet["A1"].font = Font(name="맑은 고딕", bold=True, size=16, color="111111")
    sheet["A1"].alignment = Alignment(horizontal="left", vertical="center")

    for cell in sheet[3]:
        cell.fill = PatternFill("solid", fgColor="404040")
        cell.font = Font(name="맑은 고딕", color=WHITE, size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(left=LIGHT_BORDER, right=LIGHT_BORDER)

    for row in sheet.iter_rows(min_row=4):
        sheet.row_dimensions[row[0].row].height = 22
        for cell in row:
            cell.font = Font(name="맑은 고딕", color=TEXT, size=10)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(bottom=LIGHT_BORDER)

    for column, width in {
        "A": 8,
        "B": 28,
        "C": 20,
        "D": 16,
        "E": 12,
        "F": 10,
        "G": 18,
        "H": 58,
        "I": 30,
        "J": 28,
        "K": 42,
    }.items():
        sheet.column_dimensions[column].width = width

    for row in range(4, sheet.max_row + 1):
        for column in (8, 11):
            sheet.cell(row, column).alignment = Alignment(
                horizontal="left", vertical="center", wrap_text=True
            )
        sheet.cell(row, 6).number_format = "#,##0"
        sheet.cell(row, 7).number_format = "@"
        _set_web_hyperlink(sheet.cell(row, 10))


def _add_table(sheet, name: str, ref: str | None = None, *, has_data: bool = True) -> None:
    table_ref = ref or sheet.dimensions
    if not has_data:
        sheet.auto_filter.ref = table_ref
        return
    table = Table(displayName=name, ref=table_ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)


def build_excel(
    companies: list[dict[str, Any]],
    changes: list[dict[str, Any]],
    *,
    scope: str,
    source_base_url: str,
) -> BytesIO:
    workbook = Workbook()
    companies_sheet = workbook.active
    if scope == "all":
        companies_sheet.title = "하이서울기업_참여기업 리스트"
        month_label = datetime.now(KST).strftime("%y.%m")
        _append_excel_row(
            companies_sheet,
            [f"하이서울기업 참여기업 정보({month_label})", *([None] * 10)],
        )
        _append_excel_row(companies_sheet, [None] * len(REFERENCE_COLUMNS))
        _append_excel_row(companies_sheet, [label for _, label in REFERENCE_COLUMNS])
        for sequence, company in enumerate(companies, start=1):
            _append_excel_row(
                companies_sheet,
                [
                    sequence if key == "sequence" else company.get(key)
                    for key, _ in REFERENCE_COLUMNS
                ],
        )
        _style_reference_sheet(companies_sheet)
        table_ref = f"A3:K{companies_sheet.max_row}"
        companies_sheet.auto_filter.ref = table_ref
    else:
        companies_sheet.title = "변경 기업"
        _append_excel_row(companies_sheet, [label for _, label in CHANGED_COMPANY_COLUMNS])

        for company in companies:
            values: list[Any] = []
            for key, _ in CHANGED_COMPANY_COLUMNS:
                value = company.get(key)
                if key == "last_sync_status":
                    value = STATUS_LABELS.get(str(value), value)
                elif key == "source_url" and value and str(value).startswith("/"):
                    value = f"{source_base_url.rstrip('/')}{value}"
                values.append(value)
            _append_excel_row(companies_sheet, values)

        _style_sheet(
            companies_sheet,
            {
                1: 9,
                2: 28,
                3: 18,
                4: 25,
                5: 14,
                6: 12,
                7: 11,
                8: 18,
                9: 42,
                10: 28,
                11: 30,
                12: 28,
                13: 22,
                14: 23,
                15: 23,
                16: 23,
                17: 48,
            },
        )
        for row in range(2, companies_sheet.max_row + 1):
            status_cell = companies_sheet.cell(row, 1)
            if status_cell.value == "신규":
                status_cell.fill = PatternFill("solid", fgColor=PALE_GREEN)
                status_cell.font = Font(color="1B5E20", bold=True)
            elif status_cell.value == "변경":
                status_cell.fill = PatternFill("solid", fgColor=PALE_ORANGE)
                status_cell.font = Font(color="A34A00", bold=True)
            else:
                status_cell.fill = PatternFill("solid", fgColor=PALE_GRAY)
            for column in (2, 9, 12, 13):
                companies_sheet.cell(row, column).alignment = Alignment(vertical="top", wrap_text=True)
            for column in (11, 17):
                _set_web_hyperlink(companies_sheet.cell(row, column))
            companies_sheet.cell(row, 7).number_format = "#,##0"
            companies_sheet.cell(row, 8).number_format = "@"
        _add_table(companies_sheet, "CompaniesTable", has_data=bool(companies))

    details_sheet = workbook.create_sheet("기업 상세")
    _append_excel_row(
        details_sheet,
        ["기업명", "업종", "대표자", "기업 소개", "상세 설명", "제품·서비스", "사업영역", "관심분야", "원본 페이지"]
    )
    for company in companies:
        source_url = company.get("source_url")
        if source_url and str(source_url).startswith("/"):
            source_url = f"{source_base_url.rstrip('/')}{source_url}"
        _append_excel_row(
            details_sheet,
            [
                company.get("name"),
                company.get("industry"),
                company.get("representative"),
                company.get("summary"),
                company.get("description"),
                company.get("product_service"),
                company.get("business_area"),
                company.get("interest_category"),
                source_url,
            ]
        )
    _style_sheet(
        details_sheet,
        {1: 28, 2: 18, 3: 14, 4: 48, 5: 60, 6: 48, 7: 30, 8: 22, 9: 48},
    )
    for row in range(2, details_sheet.max_row + 1):
        estimated = max(
            _estimated_lines(details_sheet.cell(row, 4).value, 42),
            _estimated_lines(details_sheet.cell(row, 5).value, 54),
            _estimated_lines(details_sheet.cell(row, 6).value, 42),
        )
        details_sheet.row_dimensions[row].height = min(180, max(42, estimated * 14))
        for column in range(4, 9):
            details_sheet.cell(row, column).alignment = Alignment(vertical="top", wrap_text=True)
        source_cell = details_sheet.cell(row, 9)
        _set_web_hyperlink(source_cell)
    _add_table(details_sheet, "CompanyDetailsTable", has_data=bool(companies))

    if scope == "changed":
        change_sheet = workbook.create_sheet("변경 이력")
        _append_excel_row(
            change_sheet,
            ["변경일시", "기업명", "기업 ID", "구분", "변경 항목", "이전 값", "새 값", "검색어"]
        )
        for change in changes:
            event_type = "신규 저장" if change.get("event_type") == "created" else "정보 변경"
            field_name = change.get("field_name")
            _append_excel_row(
                change_sheet,
                [
                    change.get("changed_at"),
                    change.get("name"),
                    change.get("source_id"),
                    event_type,
                    FIELD_LABELS.get(str(field_name), field_name or "기업 생성"),
                    change.get("old_value"),
                    change.get("new_value"),
                    change.get("query"),
                ]
            )
        _style_sheet(change_sheet, {1: 23, 2: 28, 3: 38, 4: 13, 5: 16, 6: 38, 7: 38, 8: 20})
        for row in change_sheet.iter_rows(min_row=2):
            for column in (6, 7):
                row[column - 1].alignment = Alignment(vertical="top", wrap_text=True)
            row[3].fill = PatternFill("solid", fgColor=PALE_ORANGE if row[3].value == "정보 변경" else PALE_GREEN)
        _add_table(change_sheet, "ChangeHistoryTable", has_data=bool(changes))

    info_sheet = workbook.create_sheet("내보내기 정보")
    _append_excel_row(info_sheet, ["항목", "내용"])
    _append_excel_row(info_sheet, ["생성일시", datetime.now(KST).isoformat(timespec="seconds")])
    _append_excel_row(info_sheet, ["내보내기 범위", "변경 결과" if scope == "changed" else "전체 결과"])
    _append_excel_row(info_sheet, ["기업 수", len(companies)])
    _append_excel_row(info_sheet, ["데이터 출처", f"{source_base_url.rstrip('/')}/Pages/AccountList.aspx"])
    _append_excel_row(info_sheet, ["안내", "이 파일은 로컬 데이터베이스에 저장된 최신 값을 기준으로 생성되었습니다."])
    _style_sheet(info_sheet, {1: 20, 2: 75})
    info_sheet.freeze_panes = None
    info_sheet.cell(5, 2).hyperlink = info_sheet.cell(5, 2).value
    info_sheet.cell(5, 2).style = "Hyperlink"

    workbook.properties.title = "하이서울기업 데이터 내보내기"
    workbook.properties.subject = "로컬 기업 검색·변경 이력"
    workbook.properties.creator = "하이서울기업 로컬 데이터 관리자"

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
